r"""발전소 통합목록 -> Google Sheets 업로드용 XLSX + 시트별 CSV (지침 §8).

실행: .venv\Scripts\python.exe tools\build_workbook.py

시트:
  태양광 / 육상풍력 / 해상풍력   조회용 (핵심 칼럼 우선)
  발전소통합목록                  전체 사업 대표정보
  허가이력 / 계통계약 / 출처목록 / 검수필요 / 데이터사전 / 변경이력
원천 데이터 시트에는 병합 셀을 쓰지 않는다. 첫 행 고정 + 필터 적용.
"""

from __future__ import annotations

import glob
import sys
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

from exporters.sheet_schema import (DATA_DICTIONARY, GRID_CONTRACT_COLUMNS,   # noqa: E402
                                    PERMIT_HISTORY_COLUMNS, VIEW_COLUMNS)

PROCESSED = ROOT / "data/processed"
OUT_XLSX = PROCESSED / "전남광주_발전사업_원천데이터.xlsx"
CSV_DIR = PROCESSED / "sheets_csv"

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
AUTO_FILL = PatternFill("solid", fgColor="DDEBF7")     # 자동 입력 칼럼
MANUAL_FILL = PatternFill("solid", fgColor="FFF2CC")   # 사용자 수정 칼럼(검수)


def load() -> dict[str, pd.DataFrame]:
    master = pd.read_csv(PROCESSED / "발전소_통합목록.csv", encoding="utf-8-sig")
    master = master.where(master.notna(), "")
    history = pd.read_csv(PROCESSED / "허가이력.csv", encoding="utf-8-sig").fillna("")
    contracts = pd.read_csv(PROCESSED / "계통계약.csv", encoding="utf-8-sig").fillna("")
    return {"master": master, "history": history, "contracts": contracts}


def view_sheet(master: pd.DataFrame, source: str) -> pd.DataFrame:
    subset = master[master["발전원"] == source].copy()
    subset = subset.sort_values(["시군구", "발전소명"]).reset_index(drop=True)
    subset["_row_no"] = range(1, len(subset) + 1)
    for _, origin in VIEW_COLUMNS:
        if origin not in subset.columns:
            subset[origin] = ""
    out = pd.DataFrame({label: subset[origin] for label, origin in VIEW_COLUMNS})
    return out


def source_list(master: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for _, r in master.iterrows():
        rows.append({"프로젝트ID": r["프로젝트ID"], "발전소명": r["발전소명"],
                     "대표출처": r["대표출처"], "출처페이지": r["출처페이지"],
                     "최초허가일_출처": r.get("최초허가일_출처", ""),
                     "편입경로": r.get("편입경로", "")})
    return pd.DataFrame(rows)


def review_queue() -> pd.DataFrame:
    frames = []
    for path in sorted(glob.glob(str(ROOT / "data/interim/검수필요_*.csv"))):
        item = pd.read_csv(path, encoding="utf-8-sig").fillna("")
        item.insert(0, "검수분류", Path(path).stem.replace("검수필요_", ""))
        frames.append(item)
    return pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()


def style_sheet(ws, ncols: int, manual_cols: set[int] | None = None) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    manual_cols = manual_cols or set()
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        letter = get_column_letter(c)
        width = 12
        for row in range(1, min(ws.max_row, 200) + 1):
            v = ws.cell(row=row, column=c).value
            if v is not None:
                width = max(width, min(52, len(str(v)) + 2))
        ws.column_dimensions[letter].width = width


def main() -> None:
    data = load()
    master = data["master"]
    CSV_DIR.mkdir(parents=True, exist_ok=True)

    sheets: dict[str, pd.DataFrame] = {}
    for label, source in (("태양광", "태양광"), ("육상풍력", "육상풍력"),
                          ("해상풍력", "해상풍력")):
        sheets[label] = view_sheet(master, source)
    sheets["발전소통합목록"] = master
    sheets["허가이력"] = data["history"].reindex(columns=PERMIT_HISTORY_COLUMNS)
    sheets["계통계약"] = data["contracts"].reindex(columns=GRID_CONTRACT_COLUMNS)
    sheets["출처목록"] = source_list(master)
    sheets["검수필요"] = review_queue()
    sheets["데이터사전"] = pd.DataFrame(DATA_DICTIONARY, columns=["칼럼", "설명", "예시/단위"])
    sheets["변경이력"] = pd.DataFrame(
        [["2026-09-08", "최초 구축", "회의록·허가대장·KPX 통합, 좌표 확정", "자동"]],
        columns=["일자", "구분", "내용", "작성"])

    with pd.ExcelWriter(OUT_XLSX, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            safe = frame.copy()
            safe.to_excel(writer, sheet_name=name[:31], index=False)
            style_sheet(writer.sheets[name[:31]], len(safe.columns))
            safe.to_csv(CSV_DIR / f"{name}.csv", index=False, encoding="utf-8-sig")

    print(f"XLSX: {OUT_XLSX}")
    print(f"CSV : {CSV_DIR} ({len(sheets)}개 시트)")
    for name, frame in sheets.items():
        print(f"  {name:14s} {len(frame):>5}행 x {len(frame.columns)}열")


if __name__ == "__main__":
    main()
